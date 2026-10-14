"""The SSP workbook writes text as text, never as a live formula (#218, CR-13).

openpyxl stores any string beginning with `=` as a formula, and the writer
filled cells from catalogs, the system profile, crosswalks and drafted
narratives without saying otherwise, so `=HYPERLINK(...)` or a DDE-shaped
`=cmd|...` in any of them became a formula in a file a customer opens.

The population is derived from the writer, not listed: a hostile value goes
into EVERY input `build_ssp_workbook` reads, and every cell of every sheet is
then classified. The only formulas allowed are the CIS Summary's own `=IF`
cells, whose positions come from the writer's constants. Checked twice: by
openpyxl's reader, and by the sheet XML itself, where a cell is a formula
exactly when it has an `<f>` child, which is what a spreadsheet reads.
"""

from __future__ import annotations

import zipfile
from xml.etree import ElementTree

import pytest

from policyforge.generate.policy_writer import OrgContext
from policyforge.ingest.schema import Control, ControlEnhancement
from policyforge.ssp import workbook as wb
from policyforge.ssp.narrative import SystemProfile

#: Written by hand (1d's 3597): each first character a spreadsheet reads as a
#: formula, a leading tab and carriage return, an all-formula value, and the
#: DDE and HYPERLINK shapes.
SHAPES = [
    "=1+1",
    "+1+1",
    "-1+1",
    "@SUM(1,1)",
    "\t=1+1",
    "\r=1+1",
    "=cmd|' /C calc'!A0",
    '=HYPERLINK("http://example.invalid","click")',
]
_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _build(tmp_path, shape: str):
    """A workbook with `shape` in every input the writer reads."""
    enhancement = ControlEnhancement(
        enhancement_id=f"{shape}e", title=shape, baseline=shape, description=shape
    )
    control = Control(
        control_id=f"{shape}c",
        title=shape,
        framework="NIST 800-53",
        framework_version="r5",
        family=shape,
        baseline=shape,
        control_statement=shape,
        enhancements=[enhancement],
        related_controls=[shape],
    )
    system = SystemProfile(
        **dict.fromkeys(
            (
                "name",
                "identifier",
                "owner",
                "authorizing_official",
                "security_officer",
                "operational_status",
                "system_type",
                "confidentiality",
                "integrity",
                "availability",
                "overall_categorization",
                "description",
                "environment",
                "interconnections",
            ),
            shape,
        ),
        laws_and_regulations=[shape],
    )
    out = tmp_path / "ssp.xlsx"
    wb.build_ssp_workbook(
        [control],
        system=system,
        org=OrgContext(name=shape, industry=shape),
        out_path=out,
        crosswalk={control.control_id: {"fedramp": [shape]}},
        narratives={control.control_id: shape},
        catalog_version=shape,
        generated=shape,
    )
    return out


def _cis_formula_cells() -> set[tuple[str, int, int]]:
    """Where the writer puts its own formulas, from its own constants: the
    CIS Summary's header is row 3 (a title row, a blank row), and each
    control row has one formula per status and per origination."""
    width = len(wb.IMPLEMENTATION_STATUS) + len(wb.CONTROL_ORIGINATION)
    return {(wb.SHEET_CIS, 4, column) for column in range(2, 2 + width)}


def _xml_formula_cells(path) -> set[tuple[str, int, int]]:
    """The second instrument: every cell with an `<f>` child, from the XML."""
    from openpyxl.utils.cell import coordinate_to_tuple

    found = set()
    with zipfile.ZipFile(path) as archive:
        book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        names = [s.get("name") for s in book.iter(f"{_NS}sheet")]
        for index, name in enumerate(names, start=1):
            sheet = ElementTree.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            for cell in sheet.iter(f"{_NS}c"):
                if cell.find(f"{_NS}f") is not None:
                    found.add((name, *coordinate_to_tuple(cell.get("r"))))
    return found


@pytest.mark.parametrize("shape", SHAPES, ids=repr)
def test_every_cell_filled_from_input_is_text_and_only_the_cis_cells_are_formulas(tmp_path, shape):
    from openpyxl import load_workbook

    path = _build(tmp_path, shape)
    book = load_workbook(path)
    # Every CR and CRLF is written as LF (issue 230), so a CR-led shape reads
    # back LF-led, still text, never a formula, and still quote-prefixed:
    # a leading line feed is in _FORMULA_LEADS.
    expected = shape.replace("\r", "\n")

    formulas, seeded = set(), {}
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas.add((sheet.title, cell.row, cell.column))
                elif isinstance(cell.value, str) and cell.value.startswith(expected):
                    assert cell.data_type == "s", (sheet.title, cell.coordinate)
                    assert cell.quotePrefix, (sheet.title, cell.coordinate)
                    seeded.setdefault(sheet.title, []).append(cell.value)

    assert formulas == _cis_formula_cells()
    assert _xml_formula_cells(path) == _cis_formula_cells(), "the XML agrees"
    # Extent: the seeded value reached every sheet built from input. The
    # Reference sheet's `generated` value is seeded too.
    assert set(seeded) == {
        wb.SHEET_SYSTEM,
        wb.SHEET_CONTROLS,
        wb.SHEET_ENHANCEMENTS,
        wb.SHEET_CIS,
        wb.SHEET_REFERENCE,
    }, sorted(seeded)
    # Every input value is found as given, its CRs written as LF (above).
    for value in (expected, f"{expected}c", f"{expected}e"):
        assert any(value in values for values in seeded.values()), value


def test_the_seeded_values_reach_each_sheet_in_the_numbers_the_writer_puts_them():
    """Extent, counted on one shape, so a sheet that loses a field can't hide
    behind one that kept it."""
    import tempfile
    from pathlib import Path

    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as directory:
        book = load_workbook(_build(Path(directory), "=1+1"))
        counts = {
            sheet.title: sum(
                1
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=1+1")
            )
            for sheet in book.worksheets
        }
    # System: B1 and the 14 profile fields plus laws, and the org (16 + 1).
    # Controls: id, title, family, baseline, statement, enhancements,
    # narrative, related, and the FedRAMP mapping (9).
    # Enhancements: control id, id, title, baseline, description (5).
    # CIS: the control id (1). Reference: the generated stamp (1).
    assert counts == {
        wb.SHEET_SYSTEM: 17,
        wb.SHEET_CONTROLS: 9,
        wb.SHEET_ENHANCEMENTS: 5,
        wb.SHEET_CIS: 1,
        wb.SHEET_REFERENCE: 1,
    }


# ---- passing cases ----------------------------------------------------------


def test_an_ordinary_statement_is_plain_text_without_the_quote_prefix(tmp_path):
    from openpyxl import load_workbook

    book = load_workbook(_build(tmp_path, "Accounts are reviewed quarterly."))
    cell = book[wb.SHEET_CONTROLS]["E2"]
    assert cell.value == "Accounts are reviewed quarterly."
    assert cell.data_type == "s" and not cell.quotePrefix


def test_the_cis_summary_formulas_are_still_formulas(tmp_path):
    from openpyxl import load_workbook

    book = load_workbook(_build(tmp_path, "AC-2"))
    cis = book[wb.SHEET_CIS]
    assert cis["B4"].data_type == "f"
    assert cis["B4"].value.startswith(f"=IF('{wb.SHEET_CONTROLS}'!$G2=")


def test_the_writer_has_no_numeric_column_and_a_number_would_stay_one():
    """1d's 3597: say which columns are numeric. Derived from the writer:
    every value it writes is a string, so none is. `_as_text` leaves a
    non-string alone, so a numeric column added later stays numeric."""
    from openpyxl import Workbook

    sheet = Workbook().active
    assert wb._as_text(sheet["A1"], -5).data_type == "n"
    assert wb._as_text(sheet["A2"], "-5").data_type == "s"


def test_formula_refuses_text_that_is_not_one():
    from openpyxl import Workbook

    with pytest.raises(ValueError):
        wb._formula(Workbook().active, "Implemented")
