"""ingest/parser_gate.py — S-02, model-written code checked before it runs.

The sample export is untrusted input to a prompt whose output is code, and
that code runs over the licensed file it parses. These tests are mostly
refusals, like the rest of the boundary tests, and every refusal is paired
in spirit with the ordinary parser it must not also refuse: a gate that
rejected `str.replace` would be switched off within a week.

The trial runs spawn a real child interpreter, so they are slower than the
rest of the suite. None of them touches the network — the one that tries is
refused before a packet is built, and aims at TEST-NET besides.
"""

from __future__ import annotations

import os

import pytest

from policyforge.ingest.parser_gate import check_generated_parser, trial_run
from tests.licensed_text import unmask as _m

#: A parser of the shape the codegen prompt asks for, doing nothing unusual.
ORDINARY = _m("""from __future__ import annotations

import csv
import re
from pathlib import Path

from .schema import Control


def load_@GOVRAMP_EXPORT@(export_path: Path) -> list[Control]:
    \"\"\"Reads a two-column CSV export.\"\"\"
    with open(export_path, newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    controls = []
    for row in rows:
        title = re.sub(r"\\s+", " ", row["title"]).strip().replace("  ", " ")
        controls.append(
            Control(
                control_id=row["id"],
                title=title,
                framework="GovRAMP",
                framework_version="rev5",
                source_path=str(export_path),
            )
        )
    return controls
""")


def _reasons(source: str) -> list[str]:
    return [v.reason for v in check_generated_parser(source, framework_slug="govramp")]


def _with_body(body: str, imports: str = "") -> str:
    return (
        "from __future__ import annotations\n"
        f"{imports}\n\n"
        "def load_govramp_export(export_path):\n"
        + "".join(f"    {line}\n" for line in body.splitlines())
    )


# --------------------------------------------------------------------------
# Static: what an ordinary parser is allowed
# --------------------------------------------------------------------------


def test_an_ordinary_parser_passes():
    assert _reasons(ORDINARY) == []


def test_the_hitrust_shape_passes_too():
    """Built on the public assembler, as the prompt now asks (#32)."""
    source = _with_body(
        "return build_controls([Record(reference='01.a')], source_path=str(export_path))",
        imports="from .levelled import Record, build_controls",
    )
    assert check_generated_parser(source, framework_slug="govramp") == []


def test_a_drafted_parser_may_not_build_on_the_private_hitrust_module():
    """`hitrust` moves to the private plugin (#32), so a drafted parser that
    imported it would break on every public install."""
    source = _with_body(
        "return build_controls([Record(reference='01.a')], source_path=str(export_path))",
        imports=_m("from .@HITRUST@ import Record, build_controls"),
    )
    reasons = _reasons(source)
    assert any("sibling module 'hitrust'" in r for r in reasons), reasons


def test_reading_a_workbook_passes():
    source = _with_body(
        "book = openpyxl.load_workbook(export_path, read_only=True)\n"
        "return [row for row in book.active.iter_rows(values_only=True)]",
        imports="import openpyxl",
    )
    assert _reasons(source) == []


def test_ordinary_string_and_dict_methods_are_not_mistaken_for_file_operations():
    """`replace` and `copy` are the two names a careless denylist gets wrong."""
    source = _with_body("row = {'a': 'x y'}.copy()\nreturn [row['a'].replace(' ', '')]")
    assert _reasons(source) == []


def test_open_for_reading_with_an_explicit_mode_passes():
    source = _with_body("with open(export_path, 'r', encoding='utf-8') as f:\n    return f.read()")
    assert _reasons(source) == []


# --------------------------------------------------------------------------
# Static: what it is not
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "imports",
    [
        "import requests",
        "import urllib.request",
        "from urllib import request",
        "import socket",
        "import subprocess",
        "import os",
        "import shutil",
        "import ctypes",
        "import pandas",
        "import importlib",
    ],
)
def test_an_import_outside_the_allowlist_is_refused(imports):
    reasons = _reasons(_with_body("return []", imports=imports))
    assert any("has no use for" in r for r in reasons), reasons


def test_an_import_hidden_inside_the_function_is_refused_too():
    assert any("socket" in r for r in _reasons(_with_body("import socket\nreturn []")))


def test_a_sibling_module_not_offered_is_refused():
    reasons = _reasons(_with_body("return []", imports=_m("from .@BYOC_LOADER@ import load")))
    assert any("sibling module 'byoc_loader'" in r for r in reasons)


def test_a_relative_import_reaching_outside_ingest_is_refused():
    reasons = _reasons(_with_body("return []", imports="from ..llm import base"))
    assert any("outside ingest/" in r for r in reasons)


@pytest.mark.parametrize("call", ["eval('1')", "exec('x = 1')", "__import__('os')"])
def test_turning_a_string_into_code_is_refused(call):
    assert _reasons(_with_body(f"{call}\nreturn []"))


def test_getattr_is_refused_because_it_walks_past_the_attribute_checks():
    source = _with_body(
        "from pathlib import Path\ngetattr(Path(export_path), 'write_' + 'text')('x')\nreturn []"
    )
    assert any("getattr" in r for r in _reasons(source))


def test_reaching_through_dunder_attributes_is_refused():
    reasons = _reasons(_with_body("return ().__class__.__bases__[0].__subclasses__()"))
    assert any("__class__" in r for r in reasons)
    assert any("__subclasses__" in r for r in reasons)


@pytest.mark.parametrize(
    "body",
    [
        "open(export_path, 'w').write('x')",
        "open(export_path, mode='a')",
        "open(export_path, 'r+')",
        "from pathlib import Path\nPath(export_path).open('wb')",
    ],
)
def test_opening_anything_for_writing_is_refused(body):
    reasons = _reasons(_with_body(f"{body}\nreturn []"))
    assert any("writ" in r for r in reasons), reasons


def test_a_mode_the_check_cannot_read_is_refused():
    reasons = _reasons(_with_body("mode = 'w'\nopen(export_path, mode)\nreturn []"))
    assert any("not a literal" in r for r in reasons)


@pytest.mark.parametrize(
    "body",
    [
        "from pathlib import Path\nPath('leak.txt').write_text('x')",
        "book = openpyxl.Workbook()\nbook.save('copy.xlsx')",
        "import json\njson.dump({}, None)",
    ],
)
def test_write_helpers_are_refused_by_name(body):
    reasons = _reasons(_with_body(f"{body}\nreturn []", imports="import openpyxl"))
    assert any("writes or runs something" in r for r in reasons), reasons


def test_a_parser_without_the_expected_function_is_refused():
    reasons = _reasons("def load_something_else(p):\n    return []\n")
    assert any("load_govramp_export" in r for r in reasons)


def test_code_that_is_not_python_is_refused_with_its_line():
    violations = check_generated_parser("def load_govramp_export(:\n", framework_slug="govramp")
    assert violations and "not valid Python" in violations[0].reason
    assert violations[0].line == 1


def test_every_violation_is_reported_with_a_line_number():
    source = _with_body("import socket\neval('1')\nreturn []")
    violations = check_generated_parser(source, framework_slug="govramp")
    # `_with_body` puts the function on line 4, so its body starts on 5.
    assert {v.line for v in violations} >= {5, 6}
    assert all(str(v).startswith("line ") for v in violations)


# --------------------------------------------------------------------------
# Dynamic: the trial run
# --------------------------------------------------------------------------


def _write(tmp_path, source: str, sample: str = "id,title\nAC-1,Policy\nAC-2,Accounts\n"):
    parser = tmp_path / "govramp_loader.py"
    parser.write_text(source, encoding="utf-8")
    export = tmp_path / "export.csv"
    export.write_text(sample, encoding="utf-8")
    return parser, export


def test_an_ordinary_parser_runs_and_reports_how_many_records_it_found(tmp_path):
    parser, export = _write(tmp_path, ORDINARY)
    result = trial_run(parser, export, framework_slug="govramp")

    assert result.ok, result.detail
    assert result.records == 2
    assert result.refused == ()


def test_a_parser_that_opens_a_socket_is_refused_before_it_connects(tmp_path):
    parser, export = _write(
        tmp_path,
        "def load_govramp_export(export_path):\n"
        "    import socket\n"
        "    socket.create_connection(('192.0.2.1', 80), timeout=1)\n"
        "    return []\n",
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert not result.ok
    assert any(r.startswith("socket.") for r in result.refused), result


def test_a_parser_that_writes_is_refused_and_nothing_is_written(tmp_path):
    parser, export = _write(
        tmp_path,
        "from pathlib import Path\n"
        "def load_govramp_export(export_path):\n"
        "    Path(export_path).with_suffix('.leak').write_text('licensed content')\n"
        "    return []\n",
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert not result.ok
    assert any("open for writing" in r for r in result.refused)
    assert not (tmp_path / "export.leak").exists()


def test_an_attempt_the_parser_swallows_still_fails_the_trial(tmp_path):
    """Catching the refusal and returning records anyway is still trying."""
    parser, export = _write(
        tmp_path,
        "import subprocess\n"
        "def load_govramp_export(export_path):\n"
        "    try:\n"
        "        subprocess.run(['whoami'])\n"
        "    except Exception:\n"
        "        pass\n"
        "    return [1, 2, 3]\n",
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert not result.ok
    assert any(r.startswith("subprocess.") for r in result.refused)


def test_a_parser_that_crashes_says_why(tmp_path):
    parser, export = _write(
        tmp_path, "def load_govramp_export(export_path):\n    raise ValueError('no title column')\n"
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert not result.ok
    assert "no title column" in result.detail


def test_a_parser_that_never_returns_is_stopped(tmp_path):
    parser, export = _write(
        tmp_path, "def load_govramp_export(export_path):\n    while True:\n        pass\n"
    )
    result = trial_run(parser, export, framework_slug="govramp", timeout=3)

    assert not result.ok
    assert "did not finish" in result.detail


# --------------------------------------------------------------------------
# CR-30: the interpreter's machinery reached through an allowed module, and
# reads of files the parser was not given (80's ruling on #179)
# --------------------------------------------------------------------------

#: The review's case, verbatim in shape: nothing off the allowlist is
#: imported and no forbidden builtin is named, and it passed both checks.
REVIEW_ESCAPE = "import pathlib\n" + (
    "def load_govramp_export(export_path):\n"
    "    pathlib.sys.modules['builtins'].eval('print(1)')\n"
    "    return [1]\n"
)


def test_the_reviews_escape_through_pathlib_sys_modules_is_refused():
    reasons = _reasons(REVIEW_ESCAPE)
    assert any("reaches .sys" in r for r in reasons), reasons
    assert any("reaches .modules" in r for r in reasons), reasons
    assert any("['builtins']" in r for r in reasons), reasons
    assert any("reaches .eval" in r for r in reasons), reasons


@pytest.mark.parametrize(
    "body, named",
    [
        # Each route on its own, written from what the allowed modules expose
        # (pathlib imports os and sys; csv and json import re), not from the
        # gate's own lists.
        ("import csv\ncsv.sys.modules\nreturn []", ".sys"),
        ("import pathlib\npathlib.os.getcwd()\nreturn []", ".os"),
        ("import json\nlookup = {}\nlookup['builtins']\nreturn []", "['builtins']"),
        ("import re\nlookup = {}\nlookup['__import__']('os')\nreturn []", "['__import__']"),
        ("import pathlib\nmod = pathlib.sys\nmod.modules\nreturn []", ".modules"),
        ("import io\nio.builtins\nreturn []", ".builtins"),
        ("import functools\nfunctools.reduce.__self__.exec('1')\nreturn []", ".exec"),
    ],
)
def test_each_route_to_the_machinery_is_refused_on_its_own(body, named):
    reasons = _reasons(_with_body(body))
    assert any(named in r for r in reasons), reasons


@pytest.mark.parametrize(
    "body, imports",
    [
        # What must stay: the names that sit next to the refused ones.
        ("pattern = re.compile(r'\\d+')\nreturn pattern.findall('AC-1')", "import re"),
        (
            "with Path(export_path).open(encoding='utf-8') as handle:\n"
            "    return list(csv.DictReader(handle))",
            "import csv\nfrom pathlib import Path",
        ),
        (
            "rows = list(csv.DictReader(open(export_path, encoding='utf-8')))\n"
            "return [r['open'] for r in rows] + [r['os'] for r in rows]",
            "import csv",
        ),
        (
            "book = openpyxl.load_workbook(export_path, read_only=True)\n"
            "return [row for row in book.active.iter_rows(values_only=True)]",
            "import openpyxl",
        ),
    ],
    ids=["re.compile", "Path.open on the sample", "a column called open or os", "openpyxl read"],
)
def test_the_neighbours_of_the_refused_names_still_pass(body, imports):
    assert _reasons(_with_body(body, imports=imports)) == []


@pytest.mark.parametrize(
    "source, named",
    [
        # 80's two routes (5413) and 1d's chain (on the same PR), then the two
        # string-evaluating functions ba found in the allowed modules: each
        # printed from an annotation when run.
        ("from pathlib import os\n\ndef load_govramp_export(p):\n    return os.listdir('.')\n",
         "'os'"),
        ("from pathlib import sys\n\ndef load_govramp_export(p):\n    return sys.argv\n",
         "'sys'"),
        ("import operator\nimport pathlib\n\ndef load_govramp_export(p):\n"
         "    return [operator.attrgetter('s' + 'ys')(pathlib)]\n", ".attrgetter"),
        ("import operator\nimport pathlib\n\ndef load_govramp_export(p):\n"
         "    table = operator.attrgetter('sys.modules')(pathlib)\n"
         "    builtins = operator.itemgetter('built' + 'ins')(table)\n"
         "    return [operator.attrgetter('ev' + 'al')(builtins)('1')]\n", ".attrgetter"),
        ("from operator import attrgetter\n\ndef load_govramp_export(p):\n"
         "    return [attrgetter('x')(p)]\n", "'attrgetter'"),
        ("from operator import methodcaller\n\ndef load_govramp_export(p):\n"
         "    return [methodcaller('upper')('x')]\n", "'methodcaller'"),
        ("import typing\n\ndef f(x: \"print(1) or int\"):\n    pass\n\n"
         "def load_govramp_export(p):\n    typing.get_type_hints(f)\n    return [1]\n",
         ".get_type_hints"),
        ("import functools\n\n@functools.singledispatch\ndef g(x):\n    return x\n\n"
         "def load_govramp_export(p):\n    return [g(1)]\n", ".singledispatch"),
        ("from csv import *\n\ndef load_govramp_export(p):\n    return []\n", "star import"),
        ("import pathlib\nimport string\n\ndef load_govramp_export(p):\n"
         "    f = string.Formatter()\n"
         "    table = f.get_field('0.sys.modules', (pathlib,), {})[0]\n"
         "    return [f.get_field('0.eval', (table.get('built' + 'ins'),), {})[0]]\n",
         ".Formatter"),
        ("import pathlib\nfrom string import Formatter\n\nclass Grab(Formatter):\n"
         "    def format_field(self, value, spec):\n        self.got = value\n        return ''\n\n"
         "def load_govramp_export(p):\n    g = Grab()\n    g.format('{0.sys.modules}', pathlib)\n"
         "    return [g.got]\n", "'Formatter'"),
        ("import typing\n\ndef load_govramp_export(p):\n"
         "    hint = typing.List[\"print(1) or int\"]\n"
         "    return [typing._eval_type(hint, {}, {})]\n", "._eval_type, a private name"),
        ("import typing as t\n\ndef load_govramp_export(p):\n"
         "    return [t.List._name]\n", "._name, a private name"),
        ("from typing import _eval_type\n\ndef load_govramp_export(p):\n    return []\n",
         "'_eval_type'"),
        # 1d's alias (on 5a55e4f7), and the spelling an exemption for `self`
        # would have let through: both reached eval when run (ba measured).
        ("import typing\n\ndef load_govramp_export(p):\n    t = typing\n"
         "    return [t._eval_type(typing.List['int'], {}, None)]\n",
         "._eval_type, a private name"),
        ("import typing\n\ndef load_govramp_export(p):\n    self = typing\n"
         "    return [self._eval_type(typing.List['int'], {}, None)]\n",
         "._eval_type, a private name"),
        # The stated cost: a parser's own `_x` is refused too, loudly.
        ("class Reader:\n    def __init__(self):\n        self._rows = []\n\n"
         "def load_govramp_export(p):\n    return Reader()._rows\n", "._rows, a private name"),
    ],
    ids=["from pathlib import os", "from pathlib import sys", "attrgetter of a built string",
         "1d's attrgetter chain", "attrgetter by name", "methodcaller", "get_type_hints",
         "singledispatch", "star import", "Formatter.get_field", "a Formatter subclass",
         "typing._eval_type", "a private name through an alias", "importing a private name",
         "1d's assigned alias", "an alias called self", "its own _x, the stated cost"],
)  # fmt: skip
def test_a_string_or_an_import_name_that_reaches_code_is_refused(source, named):
    reasons = _reasons(source)
    assert any(named in r for r in reasons), reasons


@pytest.mark.parametrize(
    "source",
    [
        "from typing import Optional\n\n"
        "def load_govramp_export(p) -> Optional[list]:\n    return []\n",
        "from operator import itemgetter\n\ndef load_govramp_export(p):\n"
        "    return sorted([{'id': 'b'}, {'id': 'a'}], key=itemgetter('id'))\n",
        "import dataclasses\n\n@dataclasses.dataclass\nclass Row:\n    id: 'str'\n\n"
        "def load_govramp_export(p):\n    return [Row('AC-1')]\n",
        "from functools import lru_cache\n\n@lru_cache\ndef norm(x):\n    return x.strip()\n\n"
        "def load_govramp_export(p):\n    return [norm(' AC-1 ')]\n",
        "def load_govramp_export(p):\n    return ['{0}-{1}'.format('AC', 1)]\n",
        "import string\n\ndef load_govramp_export(p):\n"
        "    return [string.Template('$id').substitute(id='AC-1'), string.capwords('access')]\n",
        "from collections import namedtuple\n\nRow = namedtuple('Row', 'id title')\n\n"
        "def load_govramp_export(p):\n    row = Row('AC-1', 'x')._replace(title='y')\n"
        "    return [row._asdict(), Row._fields, Row._make(['a', 'b'])]\n",
        "class Reader:\n    def __init__(self):\n        self.rows = []\n\n"
        "    def add(self, row):\n        self.rows.append(row)\n        return self.rows\n\n"
        "def load_govramp_export(p):\n    return Reader().add('AC-1')\n",
    ],
    ids=["typing names", "itemgetter as a sort key", "a dataclass with a string annotation",
         "lru_cache", "str.format", "string.Template and capwords",
         "namedtuple's _asdict, _replace, _fields and _make", "its own public attributes"],
)  # fmt: skip
def test_the_neighbours_of_the_refused_imports_still_pass(source):
    assert _reasons(source) == []


def _read_rule(path, *, exists):
    """The trial's read rule, as the child calls it, with `exists` given."""
    from policyforge.ingest.parser_gate import path_key, readable_path

    sample = os.path.join(os.sep, "data", "export.csv")
    root = path_key(os.path.join(os.sep, "py")).rstrip(os.sep) + os.sep
    return readable_path(path, (path_key(sample),), (root,), exists=exists)


@pytest.mark.parametrize(
    "name",
    [
        "<string>",
        "<unknown>",
        # Python 3.10, an editable install: a traceback looks up the source of
        # the finder's frame (9b on #179's PR). The whole name doesn't start
        # with "<", so a test of the whole name refused it.
        "__editable__.policyforge-1.7.0.finder.__path_hook__/<string>",
        os.path.join("somewhere", "<frozen importlib._bootstrap>"),
    ],
)
def test_a_tracebacks_source_lookup_is_not_a_read(name):
    assert _read_rule(name, exists=lambda p: False)


@pytest.mark.parametrize(
    "name, exists",
    [
        # A real file that is spelled like a lookup is still a file.
        (os.path.join(os.sep, "elsewhere", "<string>"), True),
        # Ends in ">" but its last part doesn't start with "<".
        (os.path.join(os.sep, "elsewhere", "secret>"), False),
        (os.path.join(os.sep, "elsewhere", "<dir>", "secret"), False),
        (os.path.join(os.sep, "elsewhere", "secret.txt"), True),
    ],
)
def test_anything_else_outside_the_sample_is_still_refused(name, exists):
    assert not _read_rule(name, exists=lambda p: exists)


def test_the_sample_the_roots_and_an_open_descriptor_are_readable():
    assert _read_rule(os.path.join(os.sep, "data", "export.csv"), exists=lambda p: True)
    assert _read_rule(os.path.join(os.sep, "py", "lib", "csv.py"), exists=lambda p: True)
    assert _read_rule(3, exists=lambda p: True)


def test_reading_a_file_other_than_the_sample_is_refused_in_the_trial(tmp_path):
    """The review's second case: a read the static check can't see, since
    a parser's `open` for reading is allowed and the path is only data."""
    elsewhere = tmp_path / "elsewhere.txt"
    elsewhere.write_text("not the sample\n", encoding="utf-8")
    parser, export = _write(
        tmp_path,
        "def load_govramp_export(export_path):\n"
        f"    open({str(elsewhere)!r}).read()\n"
        "    return [1]\n",
    )
    assert _reasons(parser.read_text(encoding="utf-8")) == [], "the static check can't see it"

    result = trial_run(parser, export, framework_slug="govramp")

    assert not result.ok
    assert any("read a file other than the sample" in r for r in result.refused), result
    assert any("elsewhere.txt" in r for r in result.refused), result


def test_reading_the_sample_by_path_open_runs_in_the_trial(tmp_path):
    parser, export = _write(
        tmp_path,
        "import csv\nfrom pathlib import Path\n\n"
        "def load_govramp_export(export_path):\n"
        "    with Path(export_path).open(encoding='utf-8', newline='') as handle:\n"
        "        return list(csv.DictReader(handle))\n",
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert result.ok, result.detail
    assert result.records == 2


def test_reading_a_workbook_sample_runs_in_the_trial(tmp_path):
    """openpyxl imports lazily, after the hook goes in: its own files must
    stay readable, or every workbook parser fails the trial."""
    import openpyxl

    book = openpyxl.Workbook()
    book.active.append(["id", "title"])
    book.active.append(["AC-1", "Policy"])
    export = tmp_path / "export.xlsx"
    book.save(export)
    parser = tmp_path / "govramp_loader.py"
    parser.write_text(
        "import openpyxl\n\n"
        "def load_govramp_export(export_path):\n"
        "    book = openpyxl.load_workbook(export_path, read_only=True)\n"
        "    return list(book.active.iter_rows(min_row=2, values_only=True))\n",
        encoding="utf-8",
    )
    result = trial_run(parser, export, framework_slug="govramp")

    assert result.ok, result.detail
    assert result.records == 1
